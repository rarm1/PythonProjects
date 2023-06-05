from datetime import date, datetime

import docx
import glob
import openpyxl as xl

today = datetime.now()
user_input = input("What is the date of the AOL? (Press Enter to use today's date)")
if user_input != "":
	date_format = (datetime.strptime(user_input, "%d/%m/%Y")).strftime('%A %d %B %Y')
else:
	date_format = date(day=today.day, month=today.month, year=today.year).strftime('%A %d %B %Y')
fund_list_filename = glob.glob("*.xlsx")[0]
fund_list_doc = xl.load_workbook(fund_list_filename, read_only=True, data_only=True)

print(f"Please enter the desired sheet number: ")
for idx, name in enumerate(fund_list_doc.sheetnames):
	print(f"Enter {idx} for {name}")
sheet_number = int(input())
Reader_Sheet = fund_list_doc.worksheets[sheet_number]

class Constants:
	def __init__(self):
		self.date_format = date_format
		self.AOL_Template_Doc = docx.Document("Resources/AOL_Template_4.docx")
		self.DA_Template_Doc = docx.Document("Resources/DA_Template_4.docx")
		self.fund_list_filename = fund_list_filename
		self.Reader_Sheet = Reader_Sheet
		self.FundNameColumn = None
		self.AgentCodeColumn = None
		self.ISINColumn = None
		self.DesignationColumn = None
		self.TAAddressColumn = None
		self.TAPhoneFaxColumn = None
		# Iterate through the top row of the sheet to find the column numbers of the required fields.
		for i, cell in enumerate(self.Reader_Sheet[1]):
			if cell.value == "Fund":
				self.FundNameColumn = i + 1
			elif cell.value == "Agent Code":
				self.AgentCodeColumn = i + 1
			elif cell.value == "ISIN":
				self.ISINColumn = i + 1
			elif cell.value == "Designations":
				self.DesignationColumn = i + 1
			elif cell.value == "TA Details":
				self.TAAddressColumn = i + 1
			elif cell.value == "Tel/Email:":
				self.TAPhoneFaxColumn = i + 1


